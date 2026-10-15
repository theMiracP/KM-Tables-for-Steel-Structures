"""
ГОСТ Р 21.101-2020 Form 12 / Приложение С (Титульный лист / Başlık Sayfası) Test Süiti.
Model, Konfigürasyon, CAD Table Builder, Doğrulayıcı, Docx Word Export, DXF/PDF Dışa Aktarma,
Preset/Proje Yönetimi, UI Editör Entegrasyonu ve Çok Dilli i18n testleri.
"""
import os
import pytest
from PySide6.QtCore import Qt
from PySide6.QtWidgets import QApplication

from src.models.registry import TableRegistry
from src.models.table import Table, HAlign, VAlign
from src.config.sheet_config import LayoutConfig, SheetSize, SheetOrientation, TablePlacement
from src.tables.title_sheet.config import TitleSheetConfig, DEFAULT_TITLE_SHEET_CONFIG
from src.tables.title_sheet.model import (
    TitleSheetDocument, TitleSheetSignatory, ApprovalBlock,
    build_default_title_sheet_signatories, build_empty_title_sheet_signatories
)
from src.tables.title_sheet.builder import create_title_sheet_table
from src.tables.title_sheet.validator import TitleSheetValidator
from src.tables.title_sheet.module import TitleSheetModule
from src.tables.title_sheet.editor import TitleSheetEditorWidget
from src.export.docx_exporter import DocxExporter
from src.export.dxf_exporter import DxfExporter
from src.export.pdf_exporter import PdfExporter
from src.services.preset_manager import PresetManager
from src.services.project_manager import ProjectManager
from src.ui.i18n import i18n, get_module_display_title


# =============================================================================
# 1. MODEL TESTLERİ (Приложение С / Рисунок С.1 Örneği ve Alan Bütünlüğü)
# =============================================================================
def test_title_sheet_default_sample():
    """Form 12 kurgusal/jenerik örnek verilerini doğrular (GEMINI.md Kural 6)."""
    doc = TitleSheetDocument.get_default_sample_document()

    # Kuruluş
    assert doc.organization_name == "ООО «ПРОЕКТ-СТРОЙ»"

    # Onay Bloğu
    assert doc.approval.has_approval is True
    assert doc.approval.approval_title in ("СОГЛАСОВАНО", "УТВЕРЖДАЮ")
    assert "Главный инженер" in doc.approval.role or "Директор" in doc.approval.role

    # Proje & Doküman
    assert doc.doc_code == "2026-001-КМ"
    assert "СЕВЕРНЫЙ" in doc.object_name or "Северный" in doc.object_name
    assert "ПРОЕКТНАЯ ДОКУМЕНТАЦИЯ" in doc.doc_type
    assert "Раздел" in doc.doc_title
    assert "Том 1" in doc.volume_number
    assert doc.year == "2026"

    # İmzacı Kadrosu
    assert len(doc.signatories) >= 3
    surnames = [s.surname for s in doc.signatories]
    assert any("Иванов" in s for s in surnames)
    assert any("Петров" in s for s in surnames)


def test_title_sheet_empty_document():
    """Boş şablon oluşturmayı ve en az 1 imza satırı kuralını test eder."""
    doc = TitleSheetDocument.get_empty_document()
    assert doc.higher_org == ""
    assert doc.organization_name == ""
    assert doc.doc_code == ""
    assert doc.approval.has_approval is False
    assert len(doc.signatories) >= 1


def test_title_sheet_serialization():
    """to_dict ve from_dict fonksiyonlarının veri kaybı yaşatmadığını doğrular."""
    doc = TitleSheetDocument.get_default_sample_document()
    d = doc.to_dict()

    restored = TitleSheetDocument.from_dict(d)
    assert restored.higher_org == doc.higher_org
    assert restored.organization_name == doc.organization_name
    assert restored.doc_code == doc.doc_code
    assert restored.object_name == doc.object_name
    assert restored.doc_type == doc.doc_type
    assert restored.doc_title == doc.doc_title
    assert restored.volume_number == doc.volume_number
    assert restored.year == doc.year
    assert restored.approval.has_approval == doc.approval.has_approval
    assert restored.approval.approval_title == doc.approval.approval_title
    assert len(restored.signatories) == len(doc.signatories)
    for orig_s, rest_s in zip(doc.signatories, restored.signatories):
        assert orig_s.role == rest_s.role
        assert orig_s.surname == rest_s.surname
        assert orig_s.date == rest_s.date


# =============================================================================
# 2. KONFİGÜRASYON VE BOYUT TESTLERİ (ГОСТ Р 21.101 Form 12 / Приложение Р)
# =============================================================================
def test_title_sheet_dimensions():
    """A4 çalışma alanı ve milimetrik ölçülerin GOST standartlarına uygunluğunu doğrular."""
    cfg = DEFAULT_TITLE_SHEET_CONFIG
    assert cfg.sheet_width == 210.0
    assert cfg.sheet_height == 297.0
    assert cfg.margin_left == 20.0
    assert cfg.margin_top == 5.0
    assert cfg.margin_right == 5.0
    assert cfg.margin_bottom == 5.0
    assert cfg.working_width == 185.0
    assert cfg.working_height == 287.0
    assert abs((cfg.col_role_width + cfg.col_sign_width + cfg.col_surname_width) - 185.0) < 1e-4


# =============================================================================
# 3. CAD TABLO BUILDER TESTLERİ
# =============================================================================
def test_title_sheet_builder():
    """Builder'ın 185 x 287 mm Table modelini doğru hücre ve kenarlık yapısıyla kurduğunu doğrular."""
    doc = TitleSheetDocument.get_default_sample_document()
    table = create_title_sheet_table(doc)

    assert isinstance(table, Table)
    assert abs(table.total_width - 185.0) < 1e-3
    assert abs(table.total_height - 287.0) < 1e-3
    assert len(table.columns) == 3

    # İç bölücü çizgilerin gizlendiğini (ГОСТ п. Р.2 kuralı) doğrula
    all_cells = [cell for row in table.rows for cell in row.cells]
    for c in all_cells:
        assert c.border_top is False
        assert c.border_bottom is False
        assert c.border_left is False
        assert c.border_right is False

    # Metin içeriklerinin varlığı
    texts = [c.text for c in all_cells if c.text]
    assert any("ПРОЕКТ-СТРОЙ" in t for t in texts)
    assert any("2026-001-КМ" in t for t in texts)
    assert any("2026" in t for t in texts)
    assert any("Иванов" in t for t in texts)


# =============================================================================
# 4. DOĞRULAYICI (VALIDATOR) TESTLERİ
# =============================================================================
def test_title_sheet_validator():
    """Doğrulayıcının zorunlu alan uyarı ve kontrollerini test eder."""
    validator = TitleSheetValidator()

    # Örnek belge geçerli olmalı
    sample_doc = TitleSheetDocument.get_default_sample_document()
    res = validator.validate(sample_doc)
    assert res.is_valid is True
    assert len(res.messages) == 0

    # Eksik alanlar testi
    bad_doc = TitleSheetDocument()
    bad_doc.doc_code = ""
    bad_doc.organization_name = ""
    bad_doc.doc_title = ""
    res_bad = validator.validate(bad_doc)
    assert res_bad.is_valid is False
    assert any("Belge Kodu" in msg or "Обозначение" in msg for msg in res_bad.messages)
    assert any("Kuruluş" in msg or "Организация" in msg for msg in res_bad.messages)
    assert any("bölüm" in msg or "Наименование" in msg for msg in res_bad.messages)


# =============================================================================
# 5. WORD (.DOCX) DIŞA AKTARMA TESTLERİ (DocxExporter)
# =============================================================================
def test_docx_export(tmp_path):
    """DocxExporter'ın GOST Form 12 standardında .docx dosyası ürettiğini ve içeriğini doğrular."""
    import docx

    doc = TitleSheetDocument.get_default_sample_document()
    out_file = str(tmp_path / "test_title_sheet.docx")

    exporter = DocxExporter()
    ok = exporter.export_title_sheet(doc, out_file)

    assert ok is True
    assert os.path.exists(out_file)
    assert os.path.getsize(out_file) > 2000  # En az 2 KB

    # python-docx ile içeriği doğrula
    doc_file = docx.Document(out_file)
    all_text = " ".join([p.text for p in doc_file.paragraphs])

    # Tablo metinlerini de topla
    for t in doc_file.tables:
        for r in t.rows:
            for cell in r.cells:
                all_text += " " + cell.text

    assert "ПРОЕКТ-СТРОЙ" in all_text
    assert "2026-001-КМ" in all_text
    assert "Иванов" in all_text
    assert "2026" in all_text

    # Sayfa kenar boşlukları (20, 5, 5, 5 mm)
    section = doc_file.sections[0]
    assert abs(section.left_margin.mm - 20.0) < 0.5
    assert abs(section.top_margin.mm - 5.0) < 0.5
    assert abs(section.right_margin.mm - 5.0) < 0.5
    assert abs(section.bottom_margin.mm - 5.0) < 0.5


def test_docx_export_empty_document(tmp_path):
    """Boş başlık sayfası için de çökmeden temiz .docx ürettiğini doğrular."""
    doc = TitleSheetDocument.get_empty_document()
    out_file = str(tmp_path / "empty_title_sheet.docx")

    exporter = DocxExporter()
    ok = exporter.export_title_sheet(doc, out_file)
    assert ok is True
    assert os.path.exists(out_file)
    assert os.path.getsize(out_file) > 1000


# =============================================================================
# 6. DXF VE PDF DIŞA AKTARMA TESTLERİ (1:1 Ölçek, Özel Form)
# =============================================================================
def test_title_sheet_dxf_export(tmp_path):
    """DXF çıktısının 1:1 ölçekli, Form 3 antetsiz ve metin varlığıyla üretildiğini doğrular."""
    import ezdxf

    doc = TitleSheetDocument.get_default_sample_document()
    table = create_title_sheet_table(doc)

    out_dxf = str(tmp_path / "test_title_sheet.dxf")
    exporter = DxfExporter()

    layout_cfg = LayoutConfig(
        size=SheetSize.A4,
        orientation=SheetOrientation.PORTRAIT,
        placement=TablePlacement.CENTER_SHEET
    )

    exporter.export(
        table=table,
        output_path=out_dxf,
        layout_config=layout_cfg
    )

    assert os.path.exists(out_dxf)

    # ezdxf ile doğrula
    dxf_doc = ezdxf.readfile(out_dxf)
    msp = dxf_doc.modelspace()
    assert dxf_doc.header["$INSUNITS"] == 4  # Millimeters

    texts = [e.dxf.text for e in msp.query("TEXT")]
    mtexts = [e.plain_text() for e in msp.query("MTEXT")]
    combined = " ".join(texts + mtexts)

    assert "2026-001-КМ" in combined or "ПРОЕКТ-СТРОЙ" in combined


def test_title_sheet_pdf_export(tmp_path):
    """1:1 Vektörel PDF çıktısının çökmeden üretildiğini doğrular."""
    doc = TitleSheetDocument.get_default_sample_document()
    table = create_title_sheet_table(doc)

    out_pdf = str(tmp_path / "test_title_sheet.pdf")
    exporter = PdfExporter()

    layout_cfg = LayoutConfig(
        size=SheetSize.A4,
        orientation=SheetOrientation.PORTRAIT,
        placement=TablePlacement.CENTER_SHEET
    )

    exporter.export_sheet(
        table=table,
        output_path=out_pdf,
        layout_config=layout_cfg
    )

    assert os.path.exists(out_pdf)
    assert os.path.getsize(out_pdf) > 1000


# =============================================================================
# 7. PRESET VE PROJE YÖNETİMİ TESTLERİ
# =============================================================================
def test_title_sheet_preset_manager(tmp_path):
    """PresetManager'ın Başlık Sayfası JSON şablonlarını doğru saklayıp yüklediğini doğrular."""
    pm = PresetManager.get_instance(str(tmp_path / "presets"))
    doc = TitleSheetDocument.get_default_sample_document()

    saved_path = pm.save_preset("title_sheet", "TestTitlePreset", doc, "Örnek Başlık Şablonu")
    assert os.path.exists(saved_path)

    loaded_doc = pm.load_preset("title_sheet", "TestTitlePreset")
    assert isinstance(loaded_doc, TitleSheetDocument)
    assert loaded_doc.doc_code == doc.doc_code
    assert loaded_doc.organization_name == doc.organization_name
    assert len(loaded_doc.signatories) >= 3


def test_title_sheet_project_manager(tmp_path):
    """ProjectManager'ın .kmproj dosyasında başlık sayfasını eksiksiz koruduğunu doğrular."""
    proj_path = str(tmp_path / "test_project.kmproj")
    doc = TitleSheetDocument.get_default_sample_document()

    docs = {"title_sheet": doc}
    layout_cfg = LayoutConfig(size=SheetSize.A4, orientation=SheetOrientation.PORTRAIT)
    cfgs = {"title_sheet": layout_cfg}

    ProjectManager.save_to_file(
        filepath=proj_path,
        documents=docs,
        layout_configs=cfgs,
        active_module_id="title_sheet"
    )

    assert os.path.exists(proj_path)

    restored = ProjectManager.load_from_file(proj_path)
    assert restored["active_module_id"] == "title_sheet"
    assert "title_sheet" in restored["module_documents"]
    restored_doc = restored["module_documents"]["title_sheet"]
    assert isinstance(restored_doc, TitleSheetDocument)
    assert restored_doc.doc_code == doc.doc_code
    assert restored_doc.organization_name == doc.organization_name


# =============================================================================
# 8. MODÜL VE REGİSTRY TESTLERİ
# =============================================================================
def test_title_sheet_registry():
    """TitleSheetModule'ün TableRegistry'ye doğru kayıtlı olduğunu doğrular."""
    mod = TableRegistry.get("title_sheet")
    assert mod is not None
    assert "Титульный лист" in mod.name
    assert "ГОСТ Р 21.101" in mod.description
    assert mod.supports_preview is True

    doc = mod.create_default_document()
    assert isinstance(doc, TitleSheetDocument)

    tbl = mod.build_table(doc)
    assert tbl is not None
    assert abs(tbl.total_width - 185.0) < 1e-3

    fn = mod.default_export_filename(doc)
    assert "2026-001-КМ" in fn
    assert fn.endswith(".dxf")


# =============================================================================
# 9. UI EDİTÖR WİDGET TESTLERİ
# =============================================================================
def test_title_sheet_editor_widget(qapp):
    """TitleSheetEditorWidget arayüz elemanlarının etkileşimini ve sinyallerini test eder."""
    editor = TitleSheetEditorWidget()
    doc = TitleSheetDocument.get_default_sample_document()

    changed_events = []
    editor.document_changed.connect(lambda: changed_events.append(True))

    # Doküman yükleme
    editor.set_document(doc)
    assert editor.txt_org_name.text() == "ООО «ПРОЕКТ-СТРОЙ»"
    assert editor.txt_doc_code.text() == "2026-001-КМ"
    assert editor.table_sigs.rowCount() >= 3

    # Metin değiştirme
    editor.txt_org_name.setText("ООО «ТЕСТ-ПРОЕКТ»")
    assert len(changed_events) > 0

    curr_doc = editor.get_document()
    assert curr_doc.organization_name == "ООО «ТЕСТ-ПРОЕКТ»"

    # İmza satırı ekleme
    initial_rows = editor.table_sigs.rowCount()
    editor.btn_add_sig.click()
    assert editor.table_sigs.rowCount() == initial_rows + 1

    # İmza satırı silme
    editor.table_sigs.selectRow(0)
    editor.btn_del_sig.click()
    assert editor.table_sigs.rowCount() == initial_rows

    # Örnek Şablon Butonu
    editor.btn_sample.click()
    assert editor.txt_org_name.text() == "ООО «ПРОЕКТ-СТРОЙ»"

    # Temizle (prompt=False)
    editor._clear_all_fields(prompt=False)
    assert editor.txt_org_name.text() == ""
    assert editor.table_sigs.rowCount() >= 1

    editor.close()


# =============================================================================
# 10. ÇOK DİLLİ i18n & BAŞLIK AYRIŞTIRMA TESTLERİ
# =============================================================================
def test_title_sheet_i18n():
    """Başlık sayfasının TR, RU, EN çevirilerini ve title_block ile çakışmadığını doğrular."""
    # TR
    i18n.set_language("tr")
    title_tr = get_module_display_title("title_sheet")
    assert "Başlık Sayfası" in title_tr
    assert "Form 12" in title_tr

    # RU
    i18n.set_language("ru")
    title_ru = get_module_display_title("title_sheet")
    assert "Титульный лист" in title_ru

    # EN
    i18n.set_language("en")
    title_en = get_module_display_title("title_sheet")
    assert "Title Sheet" in title_en

    # title_block ile çakışma kontrolü
    tb_title = get_module_display_title("title_block")
    assert tb_title != title_en
    assert "Form 3" in tb_title

    i18n.set_language("tr")


def test_title_sheet_editor_translations_and_combos(qapp):
    """Kırmızı kutudaki buton ve sarı kutulardaki açılır menü çevirilerini test eder."""
    i18n.set_language("tr")
    editor = TitleSheetEditorWidget()

    # 1. Kırmızı Kutudaki Buton Kontrolü:
    # "preset_btn_sample" ham anahtarı yerine "📋 Örnek Şablon" görünmeli
    assert editor.btn_sample.text() == "📋 Örnek Şablon"
    assert editor.btn_sample.text() != "preset_btn_sample"

    # 2. Sarı Kutu 1 (Onay Grifi):
    # Minimum genişlik >= 280 olmalı (metin kırpılmamalı)
    assert editor.cmb_app_title.minimumWidth() >= 280
    assert editor.cmb_app_title.count() == 2
    # Görünen metin Türkçe açıklamalı, userData ise saf GOST Rusça değeri olmalı
    assert "СОГЛАСОВАНО" in editor.cmb_app_title.itemText(0)
    assert "(Uygunluk" in editor.cmb_app_title.itemText(0)
    assert editor.cmb_app_title.itemData(0) == "СОГЛАСОВАНО"
    assert "УТВЕРЖДАЮ" in editor.cmb_app_title.itemText(1)
    assert "(Tasdik" in editor.cmb_app_title.itemText(1)
    assert editor.cmb_app_title.itemData(1) == "УТВЕРЖДАЮ"

    # 3. Sarı Kutu 2 (Doküman Türü):
    assert editor.cmb_doc_type.count() == 2
    assert "ПРОЕКТНАЯ ДОКУМЕНТАЦИЯ" in editor.cmb_doc_type.itemText(0)
    assert "Proje Dokümantasyonu" in editor.cmb_doc_type.itemText(0)
    assert editor.cmb_doc_type.itemData(0) == "ПРОЕКТНАЯ ДОКУМЕНТАЦИЯ"
    assert "РАБОЧАЯ ДОКУМЕНТАЦИЯ" in editor.cmb_doc_type.itemText(1)
    assert "Çalışma Çizimleri" in editor.cmb_doc_type.itemText(1)
    assert editor.cmb_doc_type.itemData(1) == "РАБОЧАЯ ДОКУМЕНТАЦИЯ"

    # 4. Veri Bütünlüğü (get_document saf GOST değeri döndürmeli)
    editor.cmb_app_title.setCurrentIndex(1)
    editor.cmb_doc_type.setCurrentIndex(1)
    doc = editor.get_document()
    assert doc.approval.approval_title == "УТВЕРЖДАЮ"
    assert doc.doc_type == "РАБОЧАЯ ДОКУМЕНТАЦИЯ"

    # 5. Dinamik Dil Değişimi
    i18n.set_language("en")
    assert editor.btn_sample.text() == "📋 Sample Template"
    assert "Agreed" in editor.cmb_app_title.itemText(0)
    assert "Approved" in editor.cmb_app_title.itemText(1)
    assert "Design Documentation" in editor.cmb_doc_type.itemText(0)
    assert "Working Documentation" in editor.cmb_doc_type.itemText(1)
    # Seçim korunmuş olmalı
    assert editor.cmb_app_title.currentData() == "УТВЕРЖДАЮ"
    assert editor.cmb_doc_type.currentData() == "РАБОЧАЯ ДОКУМЕНТАЦИЯ"

    # RU
    i18n.set_language("ru")
    assert editor.btn_sample.text() == "📋 Пример шаблона"
    assert "Согласующая" in editor.cmb_app_title.itemText(0)
    assert "Утверждающая" in editor.cmb_app_title.itemText(1)
    assert "Стадия «П»" in editor.cmb_doc_type.itemText(0)
    assert "Стадия «Р»" in editor.cmb_doc_type.itemText(1)

    # Geri dönüş
    i18n.set_language("tr")
    editor.close()


def test_title_sheet_editor_scroll_and_layout(qapp):
    """Başlık sayfası editörünün kaydırma alanını (QScrollArea) ve layout çakışmasızlığını doğrular."""
    from PySide6.QtWidgets import QScrollArea
    editor = TitleSheetEditorWidget()
    editor.resize(400, 500)  # Kasıtlı dar/kısa pencere boyutu
    editor.show()

    # 1. ScrollArea Kontrolü:
    assert hasattr(editor, "scroll")
    assert isinstance(editor.scroll, QScrollArea)
    assert editor.scroll.widgetResizable() is True

    # 2. Card 3 ve Card 4 elemanlarının minimum boyutları ve görünürlüğü:
    assert editor.txt_object_name.height() >= 50
    assert editor.txt_doc_title.height() >= 60
    assert editor.table_sigs.minimumHeight() >= 140

    # 3. Yön Ok Butonları (Kırpılma / 'I' görünümü engeli):
    assert editor.btn_up_sig.text() == "▲"
    assert editor.btn_down_sig.text() == "▼"

    # 4. Dinamik İmzacı Tablo Boyutlandırması:
    initial_h = editor.table_sigs.height()
    editor.btn_add_sig.click()
    assert editor.table_sigs.rowCount() > 3
    assert editor.table_sigs.height() >= initial_h

    editor.close()


def test_title_sheet_signatories_solid_delegate_no_text_overlap(qapp):
    """
    TASK-86: Form 12 Başlık Sayfası (TitleSheetEditorWidget) sorumlu imzacı kadrosunda
    veri girilirken eski metnin arkada görünmesini (%100) engelleyen SolidItemDelegate,
    opak arka plan ve hizalama mekanizmasını doğrular.
    """
    from src.ui_v2.components import SolidItemDelegate
    from PySide6.QtWidgets import QLineEdit, QStyleOptionViewItem

    editor = TitleSheetEditorWidget()
    editor.show()
    tbl = editor.table_sigs

    # 1. SolidItemDelegate atanmış olduğunu doğrula
    assert isinstance(tbl.itemDelegate(), SolidItemDelegate)

    # 2. Hücre düzenleme modunu başlat (0. satır, Unvan sütunu)
    tbl.setCurrentCell(0, 0)
    item_role = tbl.item(0, 0)
    assert item_role is not None
    assert len(item_role.text()) > 0

    tbl.editItem(item_role)
    qapp.processEvents()

    # 3. Düzenleyici QLineEdit nesnesinin opak zemin stiline sahip olduğunu doğrula
    le_role = tbl.findChild(QLineEdit)
    assert le_role is not None
    assert "background-color" in le_role.styleSheet()

    # 4. Düzenleme halindeyken hücre altındaki eski metnin boşaltıldığını doğrula (çift metin/ghosting engeli)
    opt = QStyleOptionViewItem()
    tbl.itemDelegate().initStyleOption(opt, tbl.model().index(0, 0))
    assert opt.text == ""

    # 5. Ortalanmış sütun (İmza/Paraf - sütun 1) için editörün hizalamasını doğrula
    from PySide6.QtWidgets import QAbstractItemDelegate
    tbl.closeEditor(le_role, QAbstractItemDelegate.EndEditHint.SubmitModelCache)
    qapp.processEvents()

    item_sig = tbl.item(0, 1)
    assert item_sig is not None
    tbl.setCurrentCell(0, 1)
    tbl.editItem(item_sig)
    qapp.processEvents()

    active_editors = tbl.findChildren(QLineEdit)
    assert len(active_editors) > 0
    le_sig = active_editors[-1]
    assert bool(le_sig.alignment() & Qt.AlignmentFlag.AlignHCenter)

    editor.close()
