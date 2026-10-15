"""
Word (.docx) Dışa Aktarma Motoru (DocxExporter).
ГОСТ Р 21.101-2020 Form 12 / Приложение С (Титульный лист) ve gelecekteki KM rapor tablolarını
Microsoft Word formatında üretir.
"""
import os
from typing import Optional, Any
import docx
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from src.tables.title_sheet.model import TitleSheetDocument


class DocxExporter:
    """
    ГОСТ standartlarına uygun Word (.docx) belgeleri oluşturan bağımsız dışa aktarma motoru.
    """

    def __init__(self):
        pass

    def export_title_sheet(self, document: TitleSheetDocument, output_path: str) -> bool:
        """
        TitleSheetDocument nesnesini ГОСТ Р 21.101-2020 Form 12 / Приложение С
        standardında A4 Word (.docx) belgesi olarak dışa aktarır.
        """
        try:
            doc = docx.Document()

            # 1. Pafta ve Kenar Boşlukları (A4 Dikey: 210 x 297 mm, Sol: 20 mm, Üst/Sağ/Alt: 5 mm)
            section = doc.sections[0]
            section.page_width = Mm(210.0)
            section.page_height = Mm(297.0)
            section.left_margin = Mm(20.0)
            section.right_margin = Mm(5.0)
            section.top_margin = Mm(5.0)
            section.bottom_margin = Mm(5.0)

            # 2. ГОСТ Sayfa Dış Kenarlığı (w:pgBorders)
            try:
                sectPr = section._sectPr
                # Sol 20mm, diğer kenarlar 5mm mesafede iç çerçeve
                pg_borders_xml = (
                    f'<w:pgBorders {nsdecls("w")} w:offsetFrom="page">'
                    f'<w:top w:val="single" w:sz="12" w:space="14" w:color="000000"/>'
                    f'<w:left w:val="single" w:sz="12" w:space="56" w:color="000000"/>'
                    f'<w:bottom w:val="single" w:sz="12" w:space="14" w:color="000000"/>'
                    f'<w:right w:val="single" w:sz="12" w:space="14" w:color="000000"/>'
                    f'</w:pgBorders>'
                )
                sectPr.append(parse_xml(pg_borders_xml))
            except Exception:
                pass  # Kenarlık eklenemezse içerik yine de eksiksiz üretilir

            # Varsayılan stil fontu (Arial / GOST uyumlu)
            style = doc.styles['Normal']
            font = style.font
            font.name = 'Arial'
            font.size = Pt(11)
            font.color.rgb = RGBColor(0, 0, 0)

            # İçeriğin dikeyde dengeli yayılması için boşluk yardımcısı
            def add_spacer(pt_size: int = 12):
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(pt_size)
                p.paragraph_format.line_spacing = 1.0
                return p

            # -----------------------------------------------------------------
            # 1. ÜST KURULUŞ (Поле 1)
            # -----------------------------------------------------------------
            if document.higher_org and document.higher_org.strip():
                p_high = doc.add_paragraph()
                p_high.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_high.paragraph_format.space_after = Pt(4)
                r_high = p_high.add_run(document.higher_org.strip())
                r_high.font.size = Pt(9.5)

            # -----------------------------------------------------------------
            # 2. KURULUŞ VE LOGO (Поле 2)
            # -----------------------------------------------------------------
            org_name = document.organization_name.strip()
            has_logo = document.has_logo and document.logo_image_path and os.path.exists(document.logo_image_path)

            if has_logo:
                table_logo = doc.add_table(rows=1, cols=2)
                table_logo.alignment = WD_TABLE_ALIGNMENT.CENTER
                table_logo.autofit = False
                c_logo = table_logo.cell(0, 0)
                c_text = table_logo.cell(0, 1)
                c_logo.width = Mm(30.0)
                c_text.width = Mm(150.0)

                p_img = c_logo.paragraphs[0]
                p_img.alignment = WD_ALIGN_PARAGRAPH.LEFT
                r_img = p_img.add_run()
                r_img.add_picture(document.logo_image_path, width=Mm(25.0))

                p_org = c_text.paragraphs[0]
                p_org.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r_org = p_org.add_run(org_name)
                r_org.bold = True
                r_org.font.size = Pt(12)
            else:
                p_org = doc.add_paragraph()
                p_org.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_org.paragraph_format.space_before = Pt(4)
                p_org.paragraph_format.space_after = Pt(8)
                r_org = p_org.add_run(org_name)
                r_org.bold = True
                r_org.font.size = Pt(12.5)

            # -----------------------------------------------------------------
            # 3. ONAY GRİFİ (Поле 3 — СОГЛАСОВАНО / УТВЕРЖДАЮ)
            # -----------------------------------------------------------------
            if document.approval and document.approval.has_approval:
                add_spacer(6)
                app = document.approval
                p_app = doc.add_paragraph()
                p_app.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p_app.paragraph_format.left_indent = Mm(10.0)
                p_app.paragraph_format.space_after = Pt(2)
                p_app.paragraph_format.line_spacing = 1.15

                r_app_t = p_app.add_run(f"{app.approval_title}\n")
                r_app_t.bold = True
                r_app_t.font.size = Pt(10.5)

                r_app_body = p_app.add_run(
                    f"{app.role}\n"
                    f"{app.organization}\n"
                    f"_________________ {app.name}\n"
                    f"{app.date}"
                )
                r_app_body.font.size = Pt(9.5)

            add_spacer(14)

            # -----------------------------------------------------------------
            # 4. İNŞAAT NESNESİ / PROJE ADI (Поле 4, Büyük Harflerle)
            # -----------------------------------------------------------------
            p_obj = doc.add_paragraph()
            p_obj.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_obj.paragraph_format.space_after = Pt(10)
            p_obj.paragraph_format.line_spacing = 1.2
            r_obj = p_obj.add_run(document.object_name.strip())
            r_obj.bold = True
            r_obj.font.size = Pt(13.0)

            add_spacer(8)

            # -----------------------------------------------------------------
            # 5. DOKÜMAN TÜRÜ (Поле 5: ПРОЕКТНАЯ / РАБОЧАЯ ДОКУМЕНТАЦИЯ)
            # -----------------------------------------------------------------
            p_type = doc.add_paragraph()
            p_type.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_type.paragraph_format.space_after = Pt(10)
            r_type = p_type.add_run(document.doc_type.strip())
            r_type.bold = True
            r_type.font.size = Pt(12.5)

            add_spacer(6)

            # -----------------------------------------------------------------
            # 6. DOKÜMAN / BÖLÜM ADI (Поле 6)
            # -----------------------------------------------------------------
            p_title = doc.add_paragraph()
            p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_title.paragraph_format.space_after = Pt(14)
            p_title.paragraph_format.line_spacing = 1.2
            r_title = p_title.add_run(document.doc_title.strip())
            r_title.font.size = Pt(11.0)

            add_spacer(10)

            # -----------------------------------------------------------------
            # 7. DOKÜMAN KODU & CİLT NO (Поле 7 & 8)
            # -----------------------------------------------------------------
            p_code = doc.add_paragraph()
            p_code.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_code.paragraph_format.space_after = Pt(4)
            r_code = p_code.add_run(document.doc_code.strip())
            r_code.bold = True
            r_code.font.size = Pt(14.0)

            if document.volume_number and document.volume_number.strip():
                p_vol = doc.add_paragraph()
                p_vol.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_vol.paragraph_format.space_after = Pt(14)
                r_vol = p_vol.add_run(document.volume_number.strip())
                r_vol.bold = True
                r_vol.font.size = Pt(12.0)

            add_spacer(24)

            # -----------------------------------------------------------------
            # 8. SORUMLU İMZACI KADROSU (Поля 9, 10, 11)
            # -----------------------------------------------------------------
            sigs = document.signatories or []
            if sigs:
                t_sigs = doc.add_table(rows=len(sigs), cols=3)
                t_sigs.alignment = WD_TABLE_ALIGNMENT.CENTER
                t_sigs.autofit = False

                for idx, s in enumerate(sigs):
                    row = t_sigs.rows[idx]
                    c_role = row.cells[0]
                    c_sig = row.cells[1]
                    c_name = row.cells[2]

                    c_role.width = Mm(85.0)
                    c_sig.width = Mm(40.0)
                    c_name.width = Mm(60.0)

                    # Role
                    p_r = c_role.paragraphs[0]
                    p_r.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    p_r.paragraph_format.space_after = Pt(3)
                    r_r = p_r.add_run(s.role)
                    r_r.font.size = Pt(10.0)

                    # Signature / Date
                    p_s = c_sig.paragraphs[0]
                    p_s.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_s.paragraph_format.space_after = Pt(3)
                    date_info = f" {s.date}" if s.date else ""
                    r_s = p_s.add_run(f"(подпись{date_info})")
                    r_s.font.size = Pt(8.5)
                    r_s.font.italic = True
                    r_s.font.color.rgb = RGBColor(90, 90, 90)

                    # Name
                    p_n = c_name.paragraphs[0]
                    p_n.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                    p_n.paragraph_format.space_after = Pt(3)
                    r_n = p_n.add_run(s.surname)
                    r_n.font.size = Pt(10.0)
                    r_n.bold = True

            add_spacer(18)

            # -----------------------------------------------------------------
            # 9. YIL (Поле 12, En Altta Ortalı)
            # -----------------------------------------------------------------
            p_year = doc.add_paragraph()
            p_year.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_year.paragraph_format.space_after = Pt(0)
            r_year = p_year.add_run(document.year.strip())
            r_year.bold = True
            r_year.font.size = Pt(12.0)

            # Klasör kontrolü ve kaydetme
            out_dir = os.path.dirname(output_path)
            if out_dir and not os.path.exists(out_dir):
                os.makedirs(out_dir, exist_ok=True)

            doc.save(output_path)
            return True

        except Exception as e:
            print(f"[DocxExporter] Hata: {e}")
            return False

    def export(self, document_or_table: Any, output_path: str) -> bool:
        """Genel dışa aktarma arayüzü."""
        if isinstance(document_or_table, TitleSheetDocument):
            return self.export_title_sheet(document_or_table, output_path)
        elif isinstance(document_or_table, dict):
            doc = TitleSheetDocument.from_dict(document_or_table)
            return self.export_title_sheet(doc, output_path)
        raise ValueError("DocxExporter şu anda TitleSheetDocument belgelerini desteklemektedir.")
